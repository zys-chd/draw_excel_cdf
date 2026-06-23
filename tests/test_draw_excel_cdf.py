"""reli_stat 单元测试 + 集成测试"""

import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

from draw_excel_cdf import (
    compute_statistics,
    get_limit,
    empirical_cdf,
    median_rank,
    weibull_transform,
    build_summary,
    process,
    fit_weibull,
    read_file,
)


FIXTURES = Path(__file__).parent / "fixtures"
OUTPUT_XLSX = FIXTURES / "output.xlsx"
OUTPUT_NOLIMIT_XLSX = FIXTURES / "output_nolimit.xlsx"

# ── 动态样本数据生成 ──

PARAMS_SPEC = {"Vth": (2.5, 0.2), "BVdss": (650, 10), "Rds_on": (1.8, 0.1)}


def _make_sample(n_groups=4, seed=42) -> tuple[pd.DataFrame, Path]:
    """生成随机样本数据并写入 sample.xlsx，返回 (df, path)。"""
    rng = np.random.default_rng(seed)
    # 随机 3~5 组，每组 5~15 条
    n_groups = rng.integers(3, 6) if n_groups is None else n_groups
    rows = []
    for gi in range(n_groups):
        grp_name = chr(65 + gi) + "组"  # A组, B组, C组, ...
        n = int(rng.integers(5, 16))
        for i in range(1, n + 1):
            row = {"样品编号": f"{grp_name}-{i:02d}", "批次": grp_name}
            for col, (mu, sigma) in PARAMS_SPEC.items():
                row[col] = round(float(rng.normal(mu, sigma)), 3)
            rows.append(row)

    df = pd.DataFrame(rows)
    out_path = FIXTURES / "sample.xlsx"
    FIXTURES.mkdir(parents=True, exist_ok=True)
    df.to_excel(out_path, index=False)
    return df, out_path


def _load_sample(seed=42) -> pd.DataFrame:
    df, _ = _make_sample(n_groups=None, seed=seed)
    return df


# ── empirical_cdf ──────────────────────────────


def test_empirical_cdf_single_group():
    values = pd.Series([2.0, 2.5, 2.3, 2.8, 2.1])
    cdf = empirical_cdf(values)

    assert len(cdf) == 5
    assert (cdf > 0).all()
    assert (cdf <= 1).all()
    sorted_cdf = cdf.loc[values.sort_values().index]
    assert sorted_cdf.is_monotonic_increasing


def test_empirical_cdf_ties():
    values = pd.Series([1.0, 2.0, 2.0, 3.0])
    cdf = empirical_cdf(values)
    assert cdf.iloc[1] == cdf.iloc[2]


def test_empirical_cdf_single_value():
    values = pd.Series([5.0])
    cdf = empirical_cdf(values)
    assert cdf.iloc[0] == 1.0


def test_empirical_cdf_max_n():
    values = pd.Series([1.0, 2.0, 3.0])
    cdf = empirical_cdf(values)
    assert cdf.max() == 1.0


# ── median_rank ────────────────────────────────


def test_median_rank_basic():
    values = pd.Series([2.1, 2.5, 2.3, 2.8, 2.0])
    mr = median_rank(values)
    assert len(mr) == 5
    assert (mr > 0).all()
    assert (mr < 1).all()


def test_median_rank_single():
    mr = median_rank(pd.Series([5.0]))
    assert abs(mr.iloc[0] - 0.5) < 1e-9


# ── weibull_transform ──────────────────────────


def test_weibull_basic():
    mr = pd.Series([0.1, 0.3, 0.5, 0.7, 0.9])
    w = weibull_transform(mr)
    assert w.is_monotonic_increasing
    assert not w.isna().any()


def test_weibull_boundary():
    mr = pd.Series([0.0, 0.5, 1.0])
    w = weibull_transform(mr)
    assert pd.isna(w.iloc[0])
    assert pd.isna(w.iloc[2])
    assert not pd.isna(w.iloc[1])


# ── get_limit ──────────────────────────────────


def test_get_limit_hit():
    assert get_limit("Vth", {"Vth": 3.0, "BVdss": 650}) == 3.0


def test_get_limit_regex():
    assert get_limit("Vth_25C", {r"Vth.*": 3.0}) == 3.0
    assert get_limit("Vth_150C", {r"Vth.*": 3.0}) == 3.0


def test_get_limit_first_match():
    assert get_limit("Vth_hot", {r"Vth": 3.0, r"Vth.*": 2.5}) == 3.0


def test_get_limit_miss():
    assert get_limit("Rds_on", {"Vth": 3.0}) is None


def test_get_limit_empty_map():
    assert get_limit("Vth", None) is None
    assert get_limit("Vth", {}) is None


# ── compute_statistics ─────────────────────────


def test_compute_statistics():
    df = pd.DataFrame(
        {
            "id": ["G1-1", "G1-2", "G2-1", "G2-2"],
            "group": ["A", "A", "B", "B"],
            "Vth": [2.0, 3.0, 2.5, 2.8],
        }
    )

    result = compute_statistics(df, id_col="id", group_col="group",
                                data_cols=["Vth"], limit_map={"Vth": 3.5})
    assert len(result) == 4
    assert (result["limit"] == 3.5).all()


def test_compute_statistics_limit_per_var():
    df = pd.DataFrame(
        {"id": ["X1", "X2"], "group": ["A", "A"], "Vth": [1.0, 2.0], "BV": [100, 200]}
    )
    result = compute_statistics(df, "id", "group", ["Vth", "BV"],
                                limit_map={"Vth": 3.0, "BV": 500})
    assert (result[result["variable"] == "Vth"]["limit"] == 3.0).all()
    assert (result[result["variable"] == "BV"]["limit"] == 500).all()


def test_compute_statistics_multi_col():
    df = pd.DataFrame(
        {"id": ["X1", "X2"], "group": ["A", "A"], "Vth": [1.0, 2.0], "BV": [100, 200]}
    )
    result = compute_statistics(df, "id", "group", ["Vth", "BV"])
    assert len(result) == 4


# ── build_summary ──────────────────────────────


def test_build_summary():
    df = pd.DataFrame(
        {"id": ["A1", "A2", "A3", "B1", "B2"], "group": ["G1", "G1", "G1", "G2", "G2"],
         "Vth": [2.0, 3.0, 4.0, 2.5, 2.7]}
    )
    summary = build_summary(df, "group", ["Vth"])
    assert len(summary) == 2
    assert "Weibull β" in summary.columns
    assert summary.iloc[0]["拟合 R²"] is not None


# ── fit_weibull ────────────────────────────────


def test_fit_weibull_perfect():
    """用已知 Weibull 分布生成数据，拟合应接近真实值。"""
    import numpy as np

    np.random.seed(0)
    eta_true, beta_true = 100.0, 2.5
    # 生成 Weibull 分布数据
    data = pd.Series(eta_true * (-np.log(1 - np.random.uniform(0.01, 0.99, 50))) ** (1 / beta_true))
    mr = median_rank(data)

    beta, eta, r2 = fit_weibull(data, mr)

    # beta 恢复偏差应在 30% 以内
    assert abs(beta - beta_true) / beta_true < 0.3
    # eta 恢复偏差也合理
    assert abs(eta - eta_true) / eta_true < 0.5
    assert r2 > 0.9  # 拟合度很高


def test_fit_weibull_few_points():
    """只有 2 个有效点 → 返回 None。"""
    import numpy as np
    data = pd.Series([10.0, 20.0])
    mr = pd.Series([0.3, 0.7])
    assert fit_weibull(data, mr) is None


def test_fit_weibull_with_nans():
    """数据中含 0 值时，这些点被跳过，仍能拟合。"""
    import numpy as np
    data = pd.Series([0.0, 10.0, 20.0, 30.0, 40.0])  # 0 值会被跳过
    mr = pd.Series([0.1, 0.25, 0.5, 0.75, 0.9])
    result = fit_weibull(data, mr)
    assert result is not None
    beta, eta, r2 = result
    assert beta > 0
    assert eta > 0
    assert 0 <= r2 <= 1


def test_fit_weibull_r2_range():
    """R² 应在 [0, 1] 范围内。"""
    import numpy as np
    data = pd.Series(np.random.uniform(10, 100, 30))
    mr = median_rank(data)
    result = fit_weibull(data, mr)
    assert result is not None
    _, _, r2 = result
    assert 0 <= r2 <= 1


# ── 集成测试 ───────────────────────────────────


def test_pipeline_end_to_end():
    import openpyxl
    df = _load_sample()

    output = process(
        data=df,
        id_col="样品编号", group_col="批次",
        data_cols=["Vth", "BVdss", "Rds_on"],
        output_path=OUTPUT_XLSX,
        limit_map={r"Vth": 3.0, r"BVdss": 660, r"Rds_on": 2.0},
        show_limit=True, chart_width=18, chart_height=10, marker_size=6,
    )
    assert OUTPUT_XLSX.exists()

    wb = openpyxl.load_workbook(OUTPUT_XLSX)
    assert wb.sheetnames == ["统计汇总", "Vth", "BVdss", "Rds_on"]

    ws = wb["统计汇总"]
    assert ws.cell(1, 1).value == "数据来源"
    stat_header_row = 23
    assert ws.cell(stat_header_row, 1).value == "测试项"
    assert ws.cell(stat_header_row, 15).value == "Limit"
    assert ws.cell(stat_header_row, 16).value == "limit处 CDF"
    assert ws.cell(22, 1).value == "统计汇总"
    assert ws.cell(35, 1).value == "样品编号"

    for sn in ["Vth", "BVdss", "Rds_on"]:
        ws_data = wb[sn]
        ch = ws_data._charts[0]
        assert ch.title is not None
        assert len(ch.series) >= 4, f"{sn}: {len(ch.series)} series"


def test_pipeline_no_limit():
    df = _load_sample()
    process(data=df, id_col="样品编号", group_col="批次", data_cols=["Vth"],
            output_path=OUTPUT_NOLIMIT_XLSX, show_limit=False)
    assert OUTPUT_NOLIMIT_XLSX.exists()


def test_pipeline_defaults():
    df = _load_sample()
    out = FIXTURES / "output_default.xlsx"
    process(data=df, id_col="样品编号", group_col="批次",
            data_cols=["Vth", "BVdss", "Rds_on"], output_path=out,
            chart_title="默认参数", x_label="数据", y_label="CDF")
    assert out.exists()


# ── 多尺度组合输出 ──

SCALE_COMBOS = [
    ("linear", "linear"), ("log", "linear"),
    ("linear", "log"), ("log", "log"),
]


@pytest.mark.parametrize("x_scale,y_scale", SCALE_COMBOS)
def test_scale_combos(x_scale, y_scale):
    df = _load_sample()
    out = FIXTURES / f"output_x{x_scale}_y{y_scale}.xlsx"
    process(data=df, id_col="样品编号", group_col="批次",
            data_cols=["Vth", "BVdss", "Rds_on"], output_path=out,
            limit_map={r"Vth": 3.0, r"BVdss": 660, r"Rds_on": 2.0},
            x_scale=x_scale, y_scale=y_scale, show_limit=True,
            chart_width=18, chart_height=10)
    assert out.exists()


# ── Y轴 Weibull 用例（仅 linear Y，log Y 无意义） ──

WEIBULL_SCALES = [("linear", "linear"), ("log", "linear")]


@pytest.mark.parametrize("x_scale,y_scale", WEIBULL_SCALES)
def test_weibull_yaxis(x_scale, y_scale):
    df = _load_sample()
    out = FIXTURES / f"output_weibull_x{x_scale}_y{y_scale}.xlsx"
    process(data=df, id_col="样品编号", group_col="批次",
            data_cols=["Vth"], output_path=out,
            y_axis="weibull", x_scale=x_scale, y_scale=y_scale,
            show_limit=False, chart_width=18, chart_height=10)
    assert out.exists()


# ── 图表尺寸用例 ──

CHART_SIZES = [(12, 8), (20, 12), (28, 18)]


@pytest.mark.parametrize("w,h", CHART_SIZES)
def test_chart_sizes(w, h):
    df = _load_sample()
    out = FIXTURES / f"output_size_{w}x{h}.xlsx"
    process(data=df, id_col="样品编号", group_col="批次",
            data_cols=["Vth"], output_path=out,
            chart_width=w, chart_height=h, show_limit=False)
    assert out.exists()


# ── Marker 大小用例 ──

@pytest.mark.parametrize("ms", [3, 5, 8, 12])
def test_marker_sizes(ms):
    df = _load_sample()
    out = FIXTURES / f"output_marker_{ms}.xlsx"
    process(data=df, id_col="样品编号", group_col="批次",
            data_cols=["Vth"], output_path=out,
            marker_size=ms, show_limit=False)
    assert out.exists()


# ── 自动/手动轴范围用例 ──

def test_auto_axis_on():
    df = _load_sample()
    out = FIXTURES / "output_auto_axis.xlsx"
    process(data=df, id_col="样品编号", group_col="批次",
            data_cols=["Vth"], output_path=out, auto_axis=True, show_limit=False)
    assert out.exists()


def test_auto_axis_off():
    df = _load_sample()
    out = FIXTURES / "output_manual_axis.xlsx"
    process(data=df, id_col="样品编号", group_col="批次",
            data_cols=["Vth"], output_path=out, auto_axis=False,
            x_min=2.0, x_max=3.0, y_min=0, y_max=1.1, show_limit=True)
    assert out.exists()


# ── 边界/补漏用例 ──

def test_read_file_invalid():
    """不存在的文件应抛出异常。"""
    with pytest.raises((FileNotFoundError, ValueError)):
        read_file(FIXTURES / "nonexistent.xlsx")


def test_single_group():
    """单组数据正常运行。"""
    df = pd.DataFrame({"id": ["A1", "A2", "A3"], "group": ["G1", "G1", "G1"],
                       "Vth": [2.0, 2.5, 3.0]})
    out = FIXTURES / "output_single_group.xlsx"
    process(data=df, id_col="id", group_col="group", data_cols=["Vth"],
            output_path=out, show_limit=False)
    assert out.exists()


def test_nan_in_data():
    """数据含 NaN 时不崩溃。"""
    df = pd.DataFrame({"id": ["X1", "X2", "X3", "X4"],
                       "group": ["A", "A", "B", "B"],
                       "Vth": [2.0, None, 2.5, 3.0]})
    out = FIXTURES / "output_nan_data.xlsx"
    process(data=df, id_col="id", group_col="group", data_cols=["Vth"],
            output_path=out, show_limit=False)
    assert out.exists()


def test_limit_map_partial_match():
    """limit_map 只匹配部分列。"""
    df = _load_sample()
    out = FIXTURES / "output_partial_limit.xlsx"
    process(data=df, id_col="样品编号", group_col="批次",
            data_cols=["Vth", "BVdss", "Rds_on"], output_path=out,
            limit_map={"Vth": 3.0})  # 只给 Vth 设 limit
    assert out.exists()

    import openpyxl
    wb = openpyxl.load_workbook(out)
    # Vth 有 limit 线，BVdss/Rds_on 只有组 serie
    assert len(wb["Vth"]._charts[0].series) >= 4   # 3组+limit
    assert len(wb["BVdss"]._charts[0].series) == 3  # 只有3组


def test_add_excel_chart_standalone():
    """独立使用 add_excel_chart（不经过 process）。"""
    from draw_excel_cdf import add_excel_chart
    from openpyxl import Workbook

    df = pd.DataFrame({"id": ["A1", "A2", "B1", "B2"],
                       "group": ["G1", "G1", "G2", "G2"],
                       "Vth": [1.0, 2.0, 1.5, 2.5]})
    long = compute_statistics(df, "id", "group", ["Vth"])
    wb = Workbook()
    wb = add_excel_chart(wb, long, ["Vth"],
                         x_label="Vth (V)", y_label="CDF",
                         chart_width=15, chart_height=9)
    assert len(wb["Vth"]._charts) == 1
    assert wb["Vth"]._charts[0].title is not None


def test_label_map():
    """x_label_map 正则匹配不同测试项设不同 X 轴标签。"""
    df = _load_sample()
    out = FIXTURES / "output_label_map.xlsx"
    process(data=df, id_col="样品编号", group_col="批次",
            data_cols=["Vth", "BVdss", "Rds_on"], output_path=out,
            x_label_map={r"Vth.*": "ΔVth (V)", r"BV.*": "ΔBVdss (V)", r"Rds.*": "ΔRds_on (mΩ)"},
            show_limit=False, chart_width=18, chart_height=10)
    assert out.exists()

    import openpyxl
    wb = openpyxl.load_workbook(out)
    assert "ΔVth (V)" in str(wb["Vth"]._charts[0].x_axis.title)
    assert "ΔBVdss (V)" in str(wb["BVdss"]._charts[0].x_axis.title)
    assert "ΔRds_on (mΩ)" in str(wb["Rds_on"]._charts[0].x_axis.title)


def test_cdf_values_correct():
    """验证各组 CDF 值正确。GA[1,3,2]→CDF=[0.33,1.0,0.67], GB[10,20]→CDF=[0.5,1.0]。"""
    df = pd.DataFrame({
        "id": ["A1", "A2", "A3", "B1", "B2"],
        "group": ["GA", "GA", "GA", "GB", "GB"],
        "Vth": [1.0, 3.0, 2.0, 10.0, 20.0],
    })
    out = FIXTURES / "output_cdf_check.xlsx"
    process(data=df, id_col="id", group_col="group", data_cols=["Vth"],
            output_path=out, show_limit=False)

    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws = wb["Vth"]
    rows = {}
    for r in range(2, ws.max_row + 1):
        g = ws.cell(r, 2).value
        d = ws.cell(r, 3).value
        c = ws.cell(r, 4).value
        rows.setdefault(g, []).append((d, c))

    ga = sorted(rows["GA"])
    assert len(ga) == 3
    assert abs(ga[0][1] - 1 / 3) < 0.01
    assert abs(ga[1][1] - 2 / 3) < 0.01
    assert abs(ga[2][1] - 1.0) < 0.01

    gb = sorted(rows["GB"])
    assert len(gb) == 2
    assert abs(gb[0][1] - 0.5) < 0.01
    assert abs(gb[1][1] - 1.0) < 0.01


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
