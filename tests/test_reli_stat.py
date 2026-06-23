"""reli_stat 单元测试"""

import sys
from pathlib import Path

import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

from reli_stat import (
    compute_statistics,
    get_limit,
    empirical_cdf,
    median_rank,
    weibull_transform,
    build_summary,
    process,
)


# ── empirical_cdf ──────────────────────────────


def test_empirical_cdf_single_group():
    values = pd.Series([2.0, 2.5, 2.3, 2.8, 2.1])
    cdf = empirical_cdf(values)

    assert len(cdf) == 5
    assert (cdf > 0).all()
    assert (cdf <= 1).all()
    sorted_cdf = cdf.loc[values.sort_values().index]
    assert sorted_cdf.is_monotonic_increasing


def test_empirical_cdf_ties():
    values = pd.Series([1.0, 2.0, 2.0, 3.0])
    cdf = empirical_cdf(values)
    assert cdf.iloc[1] == cdf.iloc[2]


def test_empirical_cdf_single_value():
    values = pd.Series([5.0])
    cdf = empirical_cdf(values)
    assert cdf.iloc[0] == 1.0


def test_empirical_cdf_max_n():
    values = pd.Series([1.0, 2.0, 3.0])
    cdf = empirical_cdf(values)
    assert cdf.max() == 1.0


# ── median_rank ────────────────────────────────


def test_median_rank_basic():
    values = pd.Series([2.1, 2.5, 2.3, 2.8, 2.0])
    mr = median_rank(values)
    assert len(mr) == 5
    assert (mr > 0).all()
    assert (mr < 1).all()


def test_median_rank_single():
    mr = median_rank(pd.Series([5.0]))
    assert abs(mr.iloc[0] - 0.5) < 1e-9


# ── weibull_transform ──────────────────────────


def test_weibull_basic():
    mr = pd.Series([0.1, 0.3, 0.5, 0.7, 0.9])
    w = weibull_transform(mr)
    assert w.is_monotonic_increasing
    assert not w.isna().any()


def test_weibull_boundary():
    mr = pd.Series([0.0, 0.5, 1.0])
    w = weibull_transform(mr)
    assert pd.isna(w.iloc[0])
    assert pd.isna(w.iloc[2])
    assert not pd.isna(w.iloc[1])


# ── get_limit ──────────────────────────────────


def test_get_limit_hit():
    """精确匹配测试项名。"""
    assert get_limit("Vth", {"Vth": 3.0, "BVdss": 650}) == 3.0


def test_get_limit_regex():
    """正则匹配：r\"Vth.*\" 可匹配 Vth_25C, Vth_150C 等变体。"""
    assert get_limit("Vth_25C", {r"Vth.*": 3.0}) == 3.0
    assert get_limit("Vth_150C", {r"Vth.*": 3.0}) == 3.0


def test_get_limit_first_match():
    """多个 pattern 命中时取第一个。"""
    assert get_limit("Vth_hot", {r"Vth": 3.0, r"Vth.*": 2.5}) == 3.0


def test_get_limit_miss():
    """未命中返回 None。"""
    assert get_limit("Rds_on", {"Vth": 3.0}) is None


def test_get_limit_empty_map():
    """空 map 返回 None。"""
    assert get_limit("Vth", None) is None
    assert get_limit("Vth", {}) is None


# ── compute_statistics ─────────────────────────


def test_compute_statistics():
    df = pd.DataFrame(
        {
            "id": ["G1-1", "G1-2", "G2-1", "G2-2"],
            "group": ["A", "A", "B", "B"],
            "Vth": [2.0, 3.0, 2.5, 2.8],
        }
    )

    result = compute_statistics(
        df,
        id_col="id",
        group_col="group",
        data_cols=["Vth"],
        limit_map={"Vth": 3.5},
    )

    assert list(result.columns) == [
        "id", "group", "variable", "data", "CDF", "weibull", "limit",
    ]
    assert len(result) == 4

    # Vth 的 limit 应该是 3.5（所有行一样）
    assert (result["limit"] == 3.5).all()

    # G1 的 CDF: rank=1,n=2→0.5, rank=2,n=2→1.0
    g1 = result[result["group"] == "A"].sort_values("data")
    assert g1["CDF"].iloc[0] == 0.5
    assert g1["CDF"].iloc[1] == 1.0


def test_compute_statistics_limit_per_var():
    """不同测试项可以有不同 limit。"""
    df = pd.DataFrame(
        {
            "id": ["X1", "X2"],
            "group": ["A", "A"],
            "Vth": [1.0, 2.0],
            "BV": [100, 200],
        }
    )

    result = compute_statistics(
        df, "id", "group", ["Vth", "BV"],
        limit_map={"Vth": 3.0, "BV": 500},
    )

    vth_rows = result[result["variable"] == "Vth"]
    bv_rows = result[result["variable"] == "BV"]
    assert (vth_rows["limit"] == 3.0).all()
    assert (bv_rows["limit"] == 500).all()


def test_compute_statistics_multi_col():
    df = pd.DataFrame(
        {
            "id": ["X1", "X2"],
            "group": ["A", "A"],
            "Vth": [1.0, 2.0],
            "BV": [100, 200],
        }
    )

    result = compute_statistics(df, "id", "group", ["Vth", "BV"])
    assert len(result) == 4
    assert set(result["variable"]) == {"Vth", "BV"}


# ── build_summary ──────────────────────────────


def test_build_summary():
    df = pd.DataFrame(
        {
            "id": ["A1", "A2", "A3", "B1", "B2"],
            "group": ["G1", "G1", "G1", "G2", "G2"],
            "Vth": [2.0, 3.0, 4.0, 2.5, 2.7],
        }
    )

    summary = build_summary(df, "group", ["Vth"])

    assert len(summary) == 2
    assert list(summary["测试项"]) == ["Vth", "Vth"]
    assert summary.iloc[0]["总模块数"] == 3
    assert round(summary.iloc[0]["均值"], 1) == 3.0


# ── 集成测试：端到端 pipeline ──────────────────

FIXTURES = Path(__file__).parent / "fixtures"
SAMPLE_XLSX = FIXTURES / "sample.xlsx"
OUTPUT_XLSX = FIXTURES / "output.xlsx"


def test_pipeline_end_to_end():
    """完整流水线：读 sample.xlsx → 计算 → 写 Excel → 验证输出。

    每次运行刷新 output.xlsx，可用于直接打开检查。"""
    import openpyxl

    output = process(
        input_path=SAMPLE_XLSX,
        id_col="样品编号",
        group_col="批次",
        data_cols=["Vth", "BVdss", "Rds_on"],
        output_path=OUTPUT_XLSX,
        limit_map={r"Vth": 3.0, r"BVdss": 660, r"Rds_on": 2.0},
        x_axis="data",
        y_axis="CDF",
        show_limit=True,
        chart_width=18,
        chart_height=10,
        marker_size=6,
    )
    assert output == OUTPUT_XLSX
    assert OUTPUT_XLSX.exists()

    wb = openpyxl.load_workbook(OUTPUT_XLSX)

    # --- 验证 sheet 结构 ---
    assert wb.sheetnames == ["统计汇总", "Vth", "BVdss", "Rds_on"]

    # 统计汇总
    ws = wb["统计汇总"]
    assert ws.cell(1, 1).value == "测试项"
    # 9 行统计数据 (3参数 × 3组)
    summary_rows = ws.max_row - 1 - 24 - 1  # 总行 - 统计表头 - 原始数据(24行+表头+分隔=
    # 更直接：统计数据行数 = 3参数 × 3组 = 9
    assert ws.cell(2, 1).value == "Vth"

    # 原始数据表头在统计数据后面
    # 统计数据: 1 header + 9 rows = 10，空行=11，原始表头=12
    assert ws.cell(12, 1).value == "样品编号"

    # --- 验证数据 sheet 有图表和数据 ---
    for sheet_name in ["Vth", "BVdss", "Rds_on"]:
        ws = wb[sheet_name]
        assert ws.max_row >= 2  # 有数据行
        assert len(ws._charts) == 1  # 每个 sheet 一个散点图

        chart = ws._charts[0]
        # 3 组 + limit 参考线
        assert len(chart.series) >= 3

    # --- 验证 CDF 值范围 (0, 1] ---
    ws_vth = wb["Vth"]
    cdf_values = []
    for r in range(2, ws_vth.max_row + 1):
        val = ws_vth.cell(r, 4).value  # CDF 列
        if isinstance(val, (int, float)):
            cdf_values.append(val)
    # 只要 CDF 都在 (0, 1] 且每组内单调（按 group + data 排序后验证）
    assert len(cdf_values) > 0
    assert all(0 < v <= 1 for v in cdf_values)

    # 按组验证 CDF 组内单调
    rows = []
    for r in range(2, ws_vth.max_row + 1):
        g = ws_vth.cell(r, 2).value  # group 列
        d = ws_vth.cell(r, 3).value  # data 列
        c = ws_vth.cell(r, 4).value  # CDF 列
        if all(isinstance(x, (int, float)) for x in (d, c)):
            rows.append((g, d, c))

    for grp in set(g for g, _, _ in rows):
        grp_cdf = [c for g, _, c in sorted(
            [(g, d, c) for g, d, c in rows if g == grp], key=lambda x: x[1]
        )]
        assert all(grp_cdf[i] <= grp_cdf[i + 1] for i in range(len(grp_cdf) - 1)), \
            f"CDF not monotonic within group {grp}"


def test_pipeline_no_limit():
    """不传 limit_map 时仍正常运行。"""
    output = process(
        input_path=SAMPLE_XLSX,
        id_col="样品编号",
        group_col="批次",
        data_cols=["Vth"],
        output_path=OUTPUT_XLSX,
        show_limit=False,
    )
    assert OUTPUT_XLSX.exists()
